#!/usr/bin/env python3
"""计算各客户明细中的每单毛利润和毛利率，并写回原工作簿。"""

from __future__ import annotations

import argparse
import ast
import math
import os
import re
import shutil
import sys
import tempfile
import unicodedata
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries


TOTAL_SHEETS = {"总表", "总体数据"}
RECEIPT = "收汇"
ACTUAL_COST = "实际货款"
TAX_REFUND = "退税"


class DataFormatError(Exception):
    pass


@dataclass
class Columns:
    client: int
    invoice: int | None
    date: int
    record_type: int
    description: int | None
    receipt_usd: int
    exchange_rate: int
    amount_rmb: int
    invoice_amount: int | None
    invoiced: int | None
    notes: int


@dataclass
class Block:
    title_row: int | None
    header_row: int
    first_data_row: int
    last_data_row: int
    end_row: int
    columns: Columns


def text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def normalized(value: Any) -> str:
    return re.sub(r"\s+", "", text(value))


def number(value: Any, location: str) -> float | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, bool):
        raise DataFormatError(f"{location} 不能使用TRUE/FALSE。")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        raw = text(value).replace(",", "").replace("￥", "").replace("$", "")
        negative = raw.startswith("(") and raw.endswith(")")
        if negative:
            raw = raw[1:-1]
        try:
            result = float(raw)
        except ValueError as exc:
            raise DataFormatError(f"{location} 应为数字，当前内容为“{value}”。") from exc
        if negative:
            result = -result
    if not math.isfinite(result):
        raise DataFormatError(f"{location} 不是有效数字。")
    return result


def header_columns(sheet, row: int) -> Columns | None:
    labels = {normalized(sheet.cell(row=row, column=col).value): col for col in range(1, min(sheet.max_column, 20) + 1)}

    def find(*names: str) -> int | None:
        for name in names:
            if normalized(name) in labels:
                return labels[normalized(name)]
        return None

    client = find("客户名称", "公司名称")
    date_col = find("日期")
    record_col = find("记录类型")
    usd_col = find("收汇USD")
    rate_col = find("汇率")
    amount_col = find("金额RMB", "收汇RMB")
    notes_col = find("备注")
    if None in {client, date_col, record_col, usd_col, rate_col, amount_col, notes_col}:
        return None
    return Columns(
        client=client,
        invoice=find("发票号", "订单号"),
        date=date_col,
        record_type=record_col,
        description=find("项目说明/对方", "项目说明或对方", "项目说明", "对方"),
        receipt_usd=usd_col,
        exchange_rate=rate_col,
        amount_rmb=amount_col,
        invoice_amount=find("发票金额"),
        invoiced=find("开票金额", "开票货款"),
        notes=notes_col,
    )


def is_title(sheet, row: int) -> bool:
    value = normalized(sheet.cell(row=row, column=1).value)
    return value.startswith("单据") or value.startswith("发票号")


def row_is_blank(sheet, row: int) -> bool:
    return all(text(sheet.cell(row=row, column=col).value) == "" for col in range(1, min(max(sheet.max_column, 9), 20) + 1))


def row_has_result_label(sheet, row: int) -> bool:
    labels = {normalized(sheet.cell(row=row, column=col).value) for col in range(1, min(max(sheet.max_column, 9), 20) + 1)}
    return bool(labels & {"本单毛利润", "本单毛利率"})


def find_blocks(sheet) -> list[Block]:
    headers: list[tuple[int, Columns]] = []
    for row in range(1, sheet.max_row + 1):
        columns = header_columns(sheet, row)
        if columns is not None:
            headers.append((row, columns))

    blocks: list[Block] = []
    for index, (header_row, columns) in enumerate(headers):
        next_header = headers[index + 1][0] if index + 1 < len(headers) else sheet.max_row + 1
        next_title = next((row for row in range(header_row + 1, next_header) if is_title(sheet, row)), next_header)
        end_row = min(next_header, next_title) - 1
        title_row = next((row for row in range(header_row - 1, 0, -1) if is_title(sheet, row)), None)
        result_rows = [row for row in range(header_row + 1, end_row + 1) if row_has_result_label(sheet, row)]
        data_limit = min(result_rows) - 1 if result_rows else end_row
        data_rows = [row for row in range(header_row + 1, data_limit + 1) if not row_is_blank(sheet, row)]
        if not data_rows:
            raise DataFormatError(f"{sheet.title}!A{header_row} 表头下面没有明细记录。")
        blocks.append(Block(title_row, header_row, header_row + 1, max(data_rows), end_row, columns))
    return blocks


class ValueReader:
    """读取常量、Excel公式缓存；缓存缺失时计算常见数值公式。"""

    def __init__(self, formula_workbook, cached_workbook):
        self.cached_workbook = cached_workbook
        self._merge_maps: dict[str, dict[tuple[int, int], tuple[int, int]]] = {}
        self._formula_cache: dict[tuple[str, int, int], Any] = {}
        self._formula_stack: set[tuple[str, int, int]] = set()

    def merge_anchor(self, sheet, row: int, column: int) -> tuple[int, int]:
        if sheet.title not in self._merge_maps:
            mapping: dict[tuple[int, int], tuple[int, int]] = {}
            for merged_range in sheet.merged_cells.ranges:
                min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
                for merged_row in range(min_row, max_row + 1):
                    for merged_col in range(min_col, max_col + 1):
                        mapping[(merged_row, merged_col)] = (min_row, min_col)
            self._merge_maps[sheet.title] = mapping
        return self._merge_maps[sheet.title].get((row, column), (row, column))

    def value(self, sheet, row: int, column: int) -> Any:
        row, column = self.merge_anchor(sheet, row, column)
        raw = sheet.cell(row=row, column=column).value
        if not (isinstance(raw, str) and raw.startswith("=")):
            return raw
        cached = self.cached_workbook[sheet.title].cell(row=row, column=column).value
        if cached is not None:
            return cached
        key = (sheet.title, row, column)
        if key in self._formula_cache:
            return self._formula_cache[key]
        if key in self._formula_stack:
            raise DataFormatError(f"{sheet.title}!{get_column_letter(column)}{row} 公式存在循环引用。")
        self._formula_stack.add(key)
        try:
            result = self._evaluate_formula(sheet, raw, f"{sheet.title}!{get_column_letter(column)}{row}")
            self._formula_cache[key] = result
            return result
        finally:
            self._formula_stack.remove(key)

    def number(self, sheet, row: int, column: int) -> float | None:
        location = f"{sheet.title}!{get_column_letter(column)}{row}"
        return number(self.value(sheet, row, column), location)

    def _evaluate_formula(self, sheet, formula: str, location: str) -> float:
        expression = formula[1:].strip().replace("^", "**")

        def split_args(value: str) -> list[str]:
            return [part.strip() for part in re.split(r"[,;]", value)]

        function_pattern = re.compile(r"(?i)\b(SUM|ROUND|ABS)\(([^()]*)\)")
        while True:
            match = function_pattern.search(expression)
            if not match:
                break
            function = match.group(1).upper()
            args = split_args(match.group(2))
            try:
                if function == "SUM":
                    values: list[float] = []
                    for arg in args:
                        range_match = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)", arg.upper())
                        if range_match:
                            min_col = self._column_number(range_match.group(1))
                            min_row = int(range_match.group(2))
                            max_col = self._column_number(range_match.group(3))
                            max_row = int(range_match.group(4))
                            for current_row in range(min(min_row, max_row), max(min_row, max_row) + 1):
                                for current_col in range(min(min_col, max_col), max(min_col, max_col) + 1):
                                    current = number(self.value(sheet, current_row, current_col), f"{sheet.title}!{get_column_letter(current_col)}{current_row}")
                                    if current is not None:
                                        values.append(current)
                        else:
                            values.append(self._safe_arithmetic(self._replace_references(sheet, arg), location))
                    replacement = repr(sum(values))
                elif function == "ROUND" and len(args) == 2:
                    replacement = repr(round(self._safe_arithmetic(self._replace_references(sheet, args[0]), location), int(float(args[1]))))
                elif function == "ABS" and len(args) == 1:
                    replacement = repr(abs(self._safe_arithmetic(self._replace_references(sheet, args[0]), location)))
                else:
                    raise ValueError
            except (ValueError, DataFormatError) as exc:
                if isinstance(exc, DataFormatError):
                    raise
                raise DataFormatError(f"{location} 公式暂时无法自动计算，请先用Excel打开该文件、完成计算并保存。公式：{formula}") from exc
            expression = expression[:match.start()] + replacement + expression[match.end():]

        expression = self._replace_references(sheet, expression)
        expression = re.sub(r"(?<![\w.])(\d+(?:\.\d+)?)%", r"(\1/100)", expression)
        try:
            return self._safe_arithmetic(expression, location)
        except (SyntaxError, ValueError, TypeError, ZeroDivisionError) as exc:
            raise DataFormatError(f"{location} 公式暂时无法自动计算，请先用Excel打开该文件、完成计算并保存。公式：{formula}") from exc

    def _replace_references(self, sheet, expression: str) -> str:
        pattern = re.compile(r"(?<![A-Za-z0-9_])\$?([A-Z]{1,3})\$?(\d+)(?![A-Za-z0-9_])", re.I)

        def replace(match: re.Match[str]) -> str:
            col = self._column_number(match.group(1).upper())
            row = int(match.group(2))
            value = number(self.value(sheet, row, col), f"{sheet.title}!{get_column_letter(col)}{row}")
            return "0" if value is None else repr(value)

        return pattern.sub(replace, expression)

    @staticmethod
    def _column_number(letters: str) -> int:
        result = 0
        for character in letters:
            result = result * 26 + ord(character) - 64
        return result

    @staticmethod
    def _safe_arithmetic(expression: str, location: str) -> float:
        tree = ast.parse(expression, mode="eval")
        allowed_binary = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b, ast.Mult: lambda a, b: a * b, ast.Div: lambda a, b: a / b, ast.Pow: lambda a, b: a**b}
        allowed_unary = {ast.UAdd: lambda a: a, ast.USub: lambda a: -a}

        def evaluate(node):
            if isinstance(node, ast.Expression):
                return evaluate(node.body)
            if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
                return float(node.value)
            if isinstance(node, ast.BinOp) and type(node.op) in allowed_binary:
                return allowed_binary[type(node.op)](evaluate(node.left), evaluate(node.right))
            if isinstance(node, ast.UnaryOp) and type(node.op) in allowed_unary:
                return allowed_unary[type(node.op)](evaluate(node.operand))
            raise ValueError(f"{location} 包含不支持的公式内容。")

        result = float(evaluate(tree))
        if not math.isfinite(result):
            raise ValueError(f"{location} 公式结果无效。")
        return result


def classify_record(value: Any) -> str:
    current = normalized(value)
    if not current:
        return ""
    if current == RECEIPT or "收汇" in current:
        return RECEIPT
    if current == TAX_REFUND or "退税" in current:
        return TAX_REFUND
    if current == ACTUAL_COST:
        return ACTUAL_COST
    return "其他支出"


def block_client_and_invoice(sheet, block: Block) -> tuple[str, str]:
    clients: list[str] = []
    invoices: list[str] = []
    for row in range(block.first_data_row, block.last_data_row + 1):
        client = text(sheet.cell(row=row, column=block.columns.client).value)
        if client:
            clients.append(client)
        if block.columns.invoice is not None:
            invoice = text(sheet.cell(row=row, column=block.columns.invoice).value)
            if invoice:
                invoices.append(invoice)
    unique_clients = list(dict.fromkeys(clients))
    if not unique_clients:
        raise DataFormatError(f"{sheet.title} 第{block.header_row}行开始的单据缺少客户名称。")
    if len(unique_clients) > 1:
        raise DataFormatError(f"{sheet.title} 第{block.header_row}行开始的单据区块混入了不同客户。")

    invoice = invoices[0] if invoices else ""
    if invoices and len(set(invoices)) > 1:
        raise DataFormatError(f"{sheet.title} 第{block.header_row}行开始的单据区块混入了不同发票号。")
    if not invoice and block.title_row is not None:
        title = text(sheet.cell(row=block.title_row, column=1).value)
        match = re.match(r"^\s*(?:单据\s*\d*|发票号)\s*[：:]\s*(.+?)\s*$", title)
        if match:
            invoice = match.group(1)
    if not invoice:
        raise DataFormatError(f"{sheet.title} 第{block.header_row}行开始的单据缺少发票号；请写在“发票号：实际发票号”中。")
    return unique_clients[0], invoice


def amount_column_total(sheet, block: Block, reader: ValueReader, warnings: list[str], column: int | None, label: str) -> float | None:
    if column is None:
        warnings.append(f"{sheet.title} 第{block.header_row}行开始的单据没有“{label}”列；合计保持空白。")
        return None
    seen_anchors: set[tuple[int, int]] = set()
    values: list[float] = []
    for row in range(block.first_data_row, block.last_data_row + 1):
        anchor = reader.merge_anchor(sheet, row, column)
        if anchor in seen_anchors:
            continue
        seen_anchors.add(anchor)
        current = reader.number(sheet, *anchor)
        if current is not None:
            values.append(current)
    if not values:
        client, invoice = block_client_and_invoice(sheet, block)
        warnings.append(f"单据 {client} / {invoice} 的{label}为空；{label}合计保持空白。")
        return None
    return sum(values)


def invoiced_total(sheet, block: Block, reader: ValueReader, warnings: list[str]) -> float | None:
    return amount_column_total(sheet, block, reader, warnings, block.columns.invoiced, "开票金额")


def invoice_amount_total(sheet, block: Block, reader: ValueReader, warnings: list[str]) -> float | None:
    if block.columns.invoice_amount is None:
        return None  # 旧格式没有这列，不需要额外提示。
    return amount_column_total(sheet, block, reader, warnings, block.columns.invoice_amount, "发票金额")


def calculate_block(sheet, block: Block, reader: ValueReader, warnings: list[str]) -> tuple[float | None, float | None]:
    client, invoice = block_client_and_invoice(sheet, block)
    receipt_values: list[float] = []
    actual_values: list[float] = []
    tax_values: list[float] = []
    other_values: list[float] = []
    incomplete = False

    for row in range(block.first_data_row, block.last_data_row + 1):
        if row_is_blank(sheet, row):
            continue
        record_value = sheet.cell(row=row, column=block.columns.record_type).value
        kind = classify_record(record_value)
        if not kind:
            raise DataFormatError(f"{sheet.title} 第{row}行缺少记录类型。")
        row_client = text(sheet.cell(row=row, column=block.columns.client).value)
        if row_client and row_client != client:
            raise DataFormatError(f"{sheet.title} 第{row}行客户名称与本单不一致。")
        if text(sheet.cell(row=row, column=block.columns.date).value) == "":
            warnings.append(f"{sheet.title}!{get_column_letter(block.columns.date)}{row} 日期为空；已保留空白。")

        amount = reader.number(sheet, row, block.columns.amount_rmb)
        amount_location = f"{sheet.title}!{get_column_letter(block.columns.amount_rmb)}{row}"
        if kind == RECEIPT:
            if amount is None:
                usd = reader.number(sheet, row, block.columns.receipt_usd)
                rate = reader.number(sheet, row, block.columns.exchange_rate)
                if usd is not None and rate is not None:
                    amount = usd * rate
                else:
                    warnings.append(f"{amount_location} 收汇RMB为空且无法由USD和汇率计算，本单利润保持空白。")
                    incomplete = True
            if amount is not None:
                if amount < 0:
                    raise DataFormatError(f"{amount_location} 收汇RMB应为正数。")
                receipt_values.append(amount)
        elif kind == TAX_REFUND:
            if amount is None:
                warnings.append(f"{amount_location} 退税为空，本单利润保持空白。")
                incomplete = True
            else:
                if amount < 0:
                    raise DataFormatError(f"{amount_location} 退税应为正数。")
                tax_values.append(amount)
        else:
            if amount is None:
                warnings.append(f"{amount_location} {text(record_value)}金额为空，本单利润保持空白。")
                incomplete = True
            else:
                if amount > 0:
                    raise DataFormatError(f"{amount_location} 除收汇和退税外，转出金额应填写负数。")
                if kind == ACTUAL_COST:
                    actual_values.append(amount)
                else:
                    other_values.append(amount)

    label = f"{client} / {invoice}"
    if not receipt_values:
        warnings.append(f"单据 {label} 没有完整收汇金额，本单利润保持空白。")
        incomplete = True
    if not actual_values:
        warnings.append(f"单据 {label} 没有实际货款，本单利润保持空白。")
        incomplete = True
    if not tax_values:
        warnings.append(f"单据 {label} 没有退税，本单利润保持空白。")
        incomplete = True
    if incomplete:
        return None, None

    receipt_rmb = sum(receipt_values)
    tax_refund = sum(tax_values)  # 直接使用表内退税金额，不再乘备注中的税率。
    profit = receipt_rmb + tax_refund + sum(actual_values) + sum(other_values)
    denominator = receipt_rmb + tax_refund
    margin = profit / denominator if denominator else None
    if margin is None:
        warnings.append(f"单据 {label} 的毛利率分母为0，毛利率保持空白。")
    return profit, margin


def find_label(sheet, block: Block, label: str) -> tuple[int, int] | None:
    for row in range(block.first_data_row, max(block.end_row, block.last_data_row) + 1):
        for col in range(1, min(max(sheet.max_column, 9), 20) + 1):
            if normalized(sheet.cell(row=row, column=col).value) == label:
                return row, col
    return None


def value_cell_to_right(sheet, reader: ValueReader, row: int, col: int):
    row, col = reader.merge_anchor(sheet, row, col)
    for merged_range in sheet.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= col <= merged_range.max_col:
            col = merged_range.max_col
            break
    target_row, target_col = reader.merge_anchor(sheet, row, col + 1)
    return sheet.cell(row=target_row, column=target_col)


def ensure_result_cells(sheet, block: Block, reader: ValueReader):
    profit_label = find_label(sheet, block, "本单毛利润")
    margin_label = find_label(sheet, block, "本单毛利率")
    new_style = block.columns.invoiced is not None and block.columns.invoice is None

    if new_style:
        needed = int(profit_label is None) + int(margin_label is None)
        insertion_row = block.last_data_row + 1
        if needed and insertion_row <= sheet.max_row and (is_title(sheet, insertion_row) or header_columns(sheet, insertion_row) is not None):
            sheet.insert_rows(insertion_row, needed)
        next_row = insertion_row
        if profit_label is None:
            sheet.cell(row=next_row, column=6, value="本单毛利润")
            profit_label = (next_row, 6)
            next_row += 1
        if margin_label is None:
            while profit_label and next_row == profit_label[0]:
                next_row += 1
            sheet.cell(row=next_row, column=6, value="本单毛利率")
            margin_label = (next_row, 6)
    else:
        insertion_row = block.last_data_row + 1
        if profit_label is None or margin_label is None:
            if insertion_row <= sheet.max_row and (is_title(sheet, insertion_row) or header_columns(sheet, insertion_row) is not None):
                sheet.insert_rows(insertion_row, 1)
            sheet.cell(row=insertion_row, column=1, value="本单毛利润")
            sheet.cell(row=insertion_row, column=3, value="本单毛利率")
            profit_label = (insertion_row, 1)
            margin_label = (insertion_row, 3)

    profit_cell = value_cell_to_right(sheet, reader, *profit_label)
    margin_cell = value_cell_to_right(sheet, reader, *margin_label)
    opening_total_cell = None
    opening_total_label_cell = None
    invoice_total_cell = None
    invoice_total_label_cell = None
    if new_style:
        total_row = profit_label[0]
        total_row, total_col = reader.merge_anchor(sheet, total_row, block.columns.invoiced)
        opening_total_cell = sheet.cell(row=total_row, column=total_col)
        label_row, label_col = reader.merge_anchor(sheet, total_row, block.columns.notes)
        opening_total_label_cell = sheet.cell(row=label_row, column=label_col)
        opening_total_label_cell.value = "开票金额合计"
        if block.columns.invoice_amount is not None:
            invoice_col = block.columns.invoice_amount
            if invoice_col < 2:
                raise DataFormatError(f"{sheet.title} 的“发票金额”列左侧没有位置写合计标签。")
            invoice_row, invoice_col = reader.merge_anchor(sheet, total_row, invoice_col)
            invoice_total_cell = sheet.cell(row=invoice_row, column=invoice_col)
            label_cell = sheet.cell(row=invoice_row, column=invoice_col - 1)
            if text(label_cell.value) not in {"", "发票金额合计"}:
                raise DataFormatError(f"{sheet.title}!{label_cell.coordinate} 已有内容，无法写入“发票金额合计”。")
            label_cell.value = "发票金额合计"
            invoice_total_label_cell = label_cell
    return profit_label, profit_cell, margin_label, margin_cell, opening_total_cell, opening_total_label_cell, invoice_total_cell, invoice_total_label_cell


def style_block(sheet, block: Block, profit_label, profit_cell, margin_label, margin_cell, opening_total_cell=None, opening_total_label_cell=None, invoice_total_cell=None, invoice_total_label_cell=None) -> None:
    black = Side(style="thin", color="000000")
    border = Border(left=black, right=black, top=black, bottom=black)
    result_end = max(profit_label[0], margin_label[0])
    for row in range(block.header_row, result_end + 1):
        for col in range(1, 10):
            sheet.cell(row=row, column=col).border = border
    for row, col in (profit_label, margin_label):
        label_cell = sheet.cell(row=row, column=col)
        label_font = copy(sheet.cell(row=block.header_row, column=1).font)
        label_font.bold = True
        label_cell.font = label_font
        label_cell.fill = PatternFill("solid", fgColor="FFF200")
    for result_cell in (profit_cell, margin_cell):
        result_cell.fill = PatternFill("solid", fgColor="FFF200")
    profit_cell.number_format = '#,##0.00;[Red]-#,##0.00;-'
    margin_cell.number_format = "0.0%"
    if opening_total_cell is not None and opening_total_label_cell is not None:
        total_fill = PatternFill("solid", fgColor="E2F0D9")
        opening_total_cell.fill = total_fill
        opening_total_label_cell.fill = total_fill
        total_font = copy(sheet.cell(row=block.header_row, column=1).font)
        total_font.bold = True
        opening_total_cell.font = total_font
        opening_total_label_cell.font = copy(total_font)
        opening_total_cell.number_format = '#,##0.00;[Red]-#,##0.00;-'
    if invoice_total_cell is not None and invoice_total_label_cell is not None:
        total_fill = PatternFill("solid", fgColor="E2F0D9")
        total_font = copy(sheet.cell(row=block.header_row, column=1).font)
        total_font.bold = True
        for cell in (invoice_total_cell, invoice_total_label_cell):
            cell.fill = total_fill
            cell.font = copy(total_font)
        invoice_total_cell.number_format = '#,##0.00;[Red]-#,##0.00;-'
    for col in range(1, 10):
        header_cell = sheet.cell(row=block.header_row, column=col)
        header_font = copy(header_cell.font)
        header_font.bold = True
        header_cell.font = header_font
    if block.columns.invoiced is not None:
        for row in range(block.first_data_row, block.last_data_row + 1):
            cell = sheet.cell(row=row, column=block.columns.invoiced)
            if not isinstance(cell, MergedCell):
                cell.number_format = '#,##0.00;[Red]-#,##0.00;-'
    if block.columns.invoice_amount is not None:
        for row in range(block.first_data_row, block.last_data_row + 1):
            cell = sheet.cell(row=row, column=block.columns.invoice_amount)
            if not isinstance(cell, MergedCell):
                cell.number_format = '#,##0.00;[Red]-#,##0.00;-'


def save_atomic(workbook, path: Path) -> None:
    with tempfile.NamedTemporaryFile(prefix=f".{path.stem}_", suffix=path.suffix, dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def process(input_path: Path, output_path: Path, *, backup: bool) -> tuple[int, list[str], Path | None]:
    if not input_path.exists() or input_path.suffix.lower() != ".xlsx":
        raise DataFormatError(f"找不到有效的.xlsx文件：{input_path}")
    workbook = load_workbook(input_path, data_only=False)
    cached_workbook = load_workbook(input_path, data_only=True)
    reader = ValueReader(workbook, cached_workbook)
    warnings: list[str] = []
    processed = 0

    for sheet in workbook.worksheets:
        if sheet.title in TOTAL_SHEETS:
            continue
        blocks = find_blocks(sheet)
        if not blocks:
            warnings.append(f"工作表“{sheet.title}”没有标准明细表头，已跳过。")
            continue
        calculations = [(block, *calculate_block(sheet, block, reader, warnings), invoiced_total(sheet, block, reader, warnings), invoice_amount_total(sheet, block, reader, warnings)) for block in blocks]
        for block, profit, margin, total_invoiced, total_invoice_amount in reversed(calculations):
            sheet.cell(row=block.header_row, column=block.columns.client, value="客户名称")
            if block.title_row is not None:
                _client, invoice = block_client_and_invoice(sheet, block)
                sheet.cell(row=block.title_row, column=1, value=f"发票号：{invoice}")
            profit_label, profit_cell, margin_label, margin_cell, opening_total_cell, opening_total_label_cell, invoice_total_cell, invoice_total_label_cell = ensure_result_cells(sheet, block, reader)
            profit_cell.value = profit
            margin_cell.value = margin
            if opening_total_cell is not None:
                opening_total_cell.value = total_invoiced
            if invoice_total_cell is not None:
                invoice_total_cell.value = total_invoice_amount
            style_block(sheet, block, profit_label, profit_cell, margin_label, margin_cell, opening_total_cell, opening_total_label_cell, invoice_total_cell, invoice_total_label_cell)
            processed += 1

    backup_path = None
    if input_path.resolve() == output_path.resolve() and backup:
        backup_path = input_path.with_name(f"{input_path.stem}_明细利润运行前备份{input_path.suffix}")
        shutil.copy2(input_path, backup_path)
    save_atomic(workbook, output_path)
    return processed, warnings, backup_path


def main() -> int:
    parser = argparse.ArgumentParser(description="计算各客户明细中的每单毛利润和毛利率，并直接写回工作簿。")
    parser.add_argument("input", type=Path, help="需要处理的.xlsx工作簿")
    parser.add_argument("-o", "--output", type=Path, help="测试时可指定另一个输出文件；默认直接更新输入文件")
    parser.add_argument("--no-backup", action="store_true", help="直接更新输入文件时不创建备份")
    args = parser.parse_args()
    output = args.output or args.input
    try:
        count, warnings, backup_path = process(args.input, output, backup=not args.no_backup)
    except (DataFormatError, PermissionError, OSError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    print(f"完成：已更新 {count} 个单据的毛利润和毛利率。")
    print(f"文件：{output}")
    if backup_path:
        print(f"备份：{backup_path}")
    if warnings:
        print("需要检查：")
        for warning in warnings:
            print(f"- {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
