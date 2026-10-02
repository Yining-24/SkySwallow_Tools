"""Regression coverage for Windows workbook locks and the two upload tools."""

from contextlib import closing
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from flask import Flask
from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

from skyswallow_tools.api import api_bp
from skyswallow_tools.auth import auth_bp
import tina_生成总表 as summary


FIXTURE = Path(__file__).parent / "fixtures" / "report-fixture.xlsx"


def assert_summary(content, include_ck=False):
    with closing(load_workbook(BytesIO(content), data_only=True)) as workbook:
        expected = {"总体数据", "TEST CLIENT"}
        if include_ck:
            expected.add("员工提成")
        assert set(workbook.sheetnames) == expected, workbook.sheetnames
        sheet = workbook["总体数据"]
        headers = [cell.value for cell in sheet[4]]
        profit_column = headers.index("总毛利润") + 1
        assert sheet.cell(5, profit_column).value == 3500


class ReportProcessingTests(unittest.TestCase):
    def test_validation_workbook_closes_on_success(self):
        checks = []

        def tracked_load(*args, **kwargs):
            workbook = load_workbook(*args, **kwargs)
            if kwargs.get("read_only"):
                workbook.close = Mock(wraps=workbook.close)
                checks.append(workbook)
            return workbook

        with TemporaryDirectory() as folder:
            output = Path(folder) / "summary.xlsx"
            with patch.object(summary, "load_workbook", side_effect=tracked_load):
                orders, clients, _ = summary.process(FIXTURE, output)
            self.assertEqual((orders, clients), (1, 1))
            self.assertEqual(len(checks), 1)
            checks[0].close.assert_called_once()
            assert_summary(output.read_bytes())
            # Must work immediately on Windows without forcing garbage collection.
            output.unlink()

    def test_validation_workbook_closes_on_error(self):
        checks = []

        def tracked_load(*args, **kwargs):
            workbook = load_workbook(*args, **kwargs)
            if kwargs.get("read_only"):
                workbook.close = Mock(wraps=workbook.close)
                checks.append(workbook)
                return SimpleNamespace(sheetnames=[], close=workbook.close)
            return workbook

        with TemporaryDirectory() as folder:
            output = Path(folder) / "summary.xlsx"
            with patch.object(summary, "load_workbook", side_effect=tracked_load), self.assertRaisesRegex(
                summary.DataFormatError, "sheet数量",
            ):
                summary.process(FIXTURE, output)
            self.assertEqual(len(checks), 1)
            checks[0].close.assert_called_once()
            output.unlink()

    def test_profit_then_summary_with_and_without_ck_cleans_temporary_files(self):
        app = Flask(__name__)
        stored_hash = generate_password_hash("048271", method="pbkdf2:sha256:1000")
        app.config.update(SECRET_KEY="unit-test-only", PASSWORD_HASHES={
            role: stored_hash for role in ("admin", "finance", "followup")
        })
        app.register_blueprint(auth_bp)
        app.register_blueprint(api_bp)
        app.testing = True
        created = []

        def tracked_directory(*args, **kwargs):
            temporary = TemporaryDirectory(*args, **kwargs)
            created.append(Path(temporary.name))
            return temporary

        with app.test_client() as client, patch(
            "skyswallow_tools.api.TemporaryDirectory", side_effect=tracked_directory,
        ):
            self.assertEqual(client.post("/api/auth/login", json={
                "role": "finance", "password": "048271",
            }).status_code, 200)
            profit = client.post("/api/profit", data={
                "file": (BytesIO(FIXTURE.read_bytes()), "details.xlsx"),
            })
            self.assertEqual(profit.status_code, 200, profit.data[:200])
            self.assertEqual(profit.headers["X-Processed-Blocks"], "1")
            for include_ck in (False, True, False):
                data = {"file": (BytesIO(profit.data), "profit-result.xlsx")}
                if include_ck:
                    data["ck_file"] = (BytesIO(FIXTURE.read_bytes()), "mapping.xlsx")
                response = client.post("/api/summary", data=data)
                self.assertEqual(response.status_code, 200, response.data[:200])
                self.assertEqual(response.headers["X-Order-Count"], "1")
                self.assertEqual(response.headers["X-Client-Count"], "1")
                assert_summary(response.data, include_ck)
                self.assertTrue(all(not path.exists() for path in created))


if __name__ == "__main__":
    unittest.main()
