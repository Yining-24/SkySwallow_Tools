#!/usr/bin/env python3
"""读取明细中已算好的每单利润，生成总体数据和各客户汇总sheet。"""

from __future__ import annotations

import argparse
import math
import os
import re
import sys
import tempfile
import unicodedata
from collections import OrderedDict, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from tina_明细利润 import (
    ACTUAL_COST,
    RECEIPT,
    TAX_REFUND,
    DataFormatError,
    ValueReader,
    block_client_and_invoice,
    classify_record,
    find_blocks,
    find_label,
    row_is_blank,
    text,
    value_cell_to_right,
)


TOTAL_SHEETS = {"总表", "总体数据"}
OVERVIEW_SHEET = "总体数据"
TOTAL_HEADERS = [
    "客户名称", "日期", "收汇USD", "汇率", "收汇RMB",
    "开票金额", "实际货款", "总退税", "每单毛利润", "每单毛利率",
]
OVERVIEW_HEADERS = [
    "客户名称", "收汇USD", "汇率", "收汇RMB",
    "开票金额", "实际货款", "总退税", "总毛利润", "平均毛利率",
]
COMMISSION_HEADERS = ["员工姓名", "客户类别", "提成试算RMB"]


@dataclass
class Order:
    client: str
    invoice: str
    order_date: date | None
    receipt_usd: float | None
    exchange_rate: float | None
    receipt_rmb: float | None
    invoiced: float | None
    actual_cost: float | None
    tax_refund: float | None
    gross_profit: float | None
    gross_margin: float | None
    employee: str | None = None
    category: str | None = None
    commission: float | None = None


def customer_key(value: Any) -> str:
    """客户名大小写和首尾空格不影响匹配；不模糊匹配不同名称。"""
    return unicodedata.normalize("NFKC", text(value)).casefold()


def read_ck_mapping(path: Path) -> dict[str, tuple[str, str, str | None]]:
    if not path.exists() or path.suffix.lower() != ".xlsx":
        raise DataFormatError(f"找不到有效的C/K标记.xlsx文件：{path}")
    workbook = load_workbook(path, data_only=True)
    matches = []
    for sheet in workbook.worksheets:
        for row in range(1, min(sheet.max_row, 15) + 1):
            labels = {customer_key(sheet.cell(row, col).value): col for col in range(1, min(sheet.max_column, 20) + 1)}
            name_col = labels.get("客户名") or labels.get("客户名称")
            kind_col = labels.get("标记") or labels.get("客户类别") or labels.get("c/k")
            employee_col = labels.get("员工名") or labels.get("员工姓名") or labels.get("业务员")
            if name_col and kind_col:
                matches.append((sheet, row, name_col, kind_col, employee_col))
                break
    if len(matches) != 1:
        raise DataFormatError("C/K表需有且仅有一个包含“客户名、标记”表头的sheet。")
    sheet, header_row, name_col, kind_col, employee_col = matches[0]
    mapping: dict[str, tuple[str, str, str | None]] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        name = text(sheet.cell(row, name_col).value)
        kind = text(sheet.cell(row, kind_col).value).upper()
        employee = text(sheet.cell(row, employee_col).value) if employee_col else ""
        if not name and not kind:
            continue
        if not name or kind not in {"C", "K"}:
            raise DataFormatError(f"{sheet.title} 第{row}行需要客户名，并把标记填写为C或K。")
        key = customer_key(name)
        if key in mapping:
            raise DataFormatError(f"C/K表中客户“{name}”重复。")
        mapping[key] = (name, kind, employee or None)
    if not mapping:
        raise DataFormatError("C/K表没有客户记录。")
    return mapping


def round_money(value: float) -> float:
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def apply_commissions(orders: list[Order], mapping: dict[str, tuple[str, str, str | None]], warnings: list[str]) -> None:
    used: set[str] = set()
    unmatched: set[str] = set()
    for order in orders:
        key = customer_key(order.client)
        match = mapping.get(key)
        if match is None:
            if key not in unmatched:
                warnings.append(f"客户“{order.client}”未在C/K表中；提成留空。")
                unmatched.add(key)
            continue
        used.add(key)
        order.category = match[1]
        order.employee = match[2]
        if order.gross_profit is None or order.gross_margin is None:
            warnings.append(f"{order.client} / {order.invoice} 的毛利润或毛利率为空；提成留空。")
            continue
        if order.gross_profit <= 0:
            order.commission = 0.0
        elif order.gross_margin < 0:
            warnings.append(f"{order.client} / {order.invoice} 毛利率为负但毛利润非负；提成留空，请核对。")
            continue
        else:
            reduction = min(order.gross_margin / 0.10, 1.0)
            rate = 0.05 if order.category == "C" else 0.10
            order.commission = round_money(order.gross_profit * rate * reduction)
    extra = [name for key, (name, _, _) in mapping.items() if key not in used]
    if extra:
        warnings.append("C/K表中未出现在明细里的客户：" + "、".join(extra))
    if any(order.category == "K" for order in orders):
        warnings.append("K类客户的12个月保护期缺起算日期；金额仅供试算，不能视为应发提成。")
    if any(order.order_date is not None and order.order_date.year != 2027 for order in orders):
        warnings.append("明细含非2027年单据；这些提成金额只是套用2027制度的演示。")


def number(value: Any, location: str) -> float | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, bool):
        raise DataFormatError(f"{location} 不能使用TRUE/FALSE。")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        raw = unicodedata.normalize("NFKC", str(value)).strip().replace(",", "").replace("￥", "").replace("$", "")
        negative = raw.startswith("(") and raw.endswith(")")
        if negative:
            raw = raw[1:-1]
        try:
            result = float(raw)
        except ValueError as exc:
            raise DataFormatError(f"{location} 应为数字，当前内容为“{value}”。") from exc
        if negative:
            result = -result
    if not math.isfinite(result):
        raise DataFormatError(f"{location} 不是有效数字。")
    return result


def parsed_date(value: Any, location: str) -> date | None:
    if value is None or text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        return converted.date() if isinstance(converted, datetime) else converted
    numeric_text = text(value).replace(",", "")
    if re.fullmatch(r"\d+(?:\.\d+)?", numeric_text):
        converted = from_excel(float(numeric_text))
        return converted.date() if isinstance(converted, datetime) else converted
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%m-%d-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text(value), fmt).date()
        except ValueError:
            pass
    raise DataFormatError(f"{location} 日期格式无法识别：{value}")


def complete_sum(values: list[float | None], *, require_rows: bool = True) -> float | None:
    if require_rows and not values:
        return None
    if any(value is None for value in values):
        return None
    return sum(value for value in values if value is not None)


def read_block(sheet, block, reader: ValueReader, warnings: list[str]) -> Order:
    client, invoice = block_client_and_invoice(sheet, block)
    records: dict[str, list[tuple[int, date | None, float | None, float | None, float | None]]] = defaultdict(list)

    for row in range(block.first_data_row, block.last_data_row + 1):
        if row_is_blank(sheet, row):
            continue
        record_value = sheet.cell(row=row, column=block.columns.record_type).value
        record_kind = classify_record(record_value)
        if not record_kind:
            raise DataFormatError(f"{sheet.title} 第{row}行缺少记录类型。")
        row_client = text(sheet.cell(row=row, column=block.columns.client).value)
        if row_client and row_client != client:
            raise DataFormatError(f"{sheet.title} 第{row}行客户名称与本单不一致。")

        date_value = reader.value(sheet, row, block.columns.date)
        date_location = f"{sheet.title}!{get_column_letter(block.columns.date)}{row}"
        record_date = parsed_date(date_value, date_location)
        if record_date is None:
            warnings.append(f"{date_location} 日期为空；总表不会为该空白补日期。")
        usd = reader.number(sheet, row, block.columns.receipt_usd)
        rate = reader.number(sheet, row, block.columns.exchange_rate)
        amount = reader.number(sheet, row, block.columns.amount_rmb)
        amount_location = f"{sheet.title}!{get_column_letter(block.columns.amount_rmb)}{row}"
        if amount is not None:
            if record_kind in {RECEIPT, TAX_REFUND} and amount < 0:
                raise DataFormatError(f"{amount_location} {record_kind}应为正数。")
            if record_kind not in {RECEIPT, TAX_REFUND} and amount > 0:
                raise DataFormatError(f"{amount_location} 除收汇和退税外，转出金额应填写负数。")
        records[record_kind].append((row, record_date, usd, rate, amount))

    receipt_rows = records[RECEIPT]
    if not receipt_rows or any(record[1] is None for record in receipt_rows):
        order_date = None
    else:
        order_date = max(record[1] for record in receipt_rows if record[1] is not None)

    receipt_usd = complete_sum([record[2] for record in receipt_rows])
    if receipt_usd is None:
        warnings.append(f"单据 {client} / {invoice} 的收汇USD不完整；总表收汇USD和汇率保持空白。")

    rmb_values: list[float | None] = []
    for row, _record_date, usd, rate, amount in receipt_rows:
        if amount is not None:
            rmb_values.append(amount)
        elif usd is not None and rate is not None:
            rmb_values.append(usd * rate)
        else:
            rmb_values.append(None)
            warnings.append(f"{sheet.title}!{get_column_letter(block.columns.amount_rmb)}{row} 收汇RMB为空且无法计算；总表保持空白。")
    receipt_rmb = complete_sum(rmb_values)
    if receipt_rmb is None:
        warnings.append(f"单据 {client} / {invoice} 的收汇RMB不完整；总表收汇RMB和汇率保持空白。")
    exchange_rate = receipt_rmb / receipt_usd if receipt_rmb is not None and receipt_usd not in (None, 0) else None

    def sum_type(record_kind: str, label: str) -> float | None:
        values = [record[4] for record in records[record_kind]]
        result = complete_sum(values)
        if result is None:
            warnings.append(f"单据 {client} / {invoice} 的{label}不完整；总表对应单元格保持空白。")
        return result

    actual_cost = sum_type(ACTUAL_COST, "实际货款")
    tax_refund = sum_type(TAX_REFUND, "退税")  # 表中金额就是退税额，不再按备注税率计算。

    invoiced = None
    if block.columns.invoiced is None:
        warnings.append(f"单据 {client} / {invoice} 没有“开票金额”列；总表保持空白。")
    else:
        seen_anchors: set[tuple[int, int]] = set()
        invoice_values: list[float] = []
        for row in range(block.first_data_row, block.last_data_row + 1):
            anchor = reader.merge_anchor(sheet, row, block.columns.invoiced)
            if anchor in seen_anchors:
                continue
            seen_anchors.add(anchor)
            current = reader.number(sheet, *anchor)
            if current is not None:
                invoice_values.append(current)
        if invoice_values:
            invoiced = sum(invoice_values)
        else:
            warnings.append(f"单据 {client} / {invoice} 的开票金额为空；总表保持空白。")

    gross_profit = None
    gross_margin = None
    profit_label = find_label(sheet, block, "本单毛利润")
    margin_label = find_label(sheet, block, "本单毛利率")
    if profit_label is None:
        warnings.append(f"单据 {client} / {invoice} 没有本单毛利润；总表保持空白。")
    else:
        profit_cell = value_cell_to_right(sheet, reader, *profit_label)
        gross_profit = reader.number(sheet, profit_cell.row, profit_cell.column)
        if gross_profit is None:
            warnings.append(f"{sheet.title}!{profit_cell.coordinate} 本单毛利润为空；总表保持空白。")
    if margin_label is None:
        warnings.append(f"单据 {client} / {invoice} 没有本单毛利率；总表保持空白。")
    else:
        margin_cell = value_cell_to_right(sheet, reader, *margin_label)
        gross_margin = reader.number(sheet, margin_cell.row, margin_cell.column)
        if gross_margin is None:
            warnings.append(f"{sheet.title}!{margin_cell.coordinate} 本单毛利率为空；总表保持空白。")

    return Order(client, invoice, order_date, receipt_usd, exchange_rate, receipt_rmb, invoiced, actual_cost, tax_refund, gross_profit, gross_margin)


def aggregate(orders: list[Order], *, client_label: str | None = None, invoice_label: str = "客户合计") -> Order:
    def field_sum(name: str) -> float | None:
        return complete_sum([getattr(order, name) for order in orders])

    receipt_usd = field_sum("receipt_usd")
    receipt_rmb = field_sum("receipt_rmb")
    tax_refund = field_sum("tax_refund")
    gross_profit = field_sum("gross_profit")
    exchange_rate = receipt_rmb / receipt_usd if receipt_rmb is not None and receipt_usd not in (None, 0) else None
    denominator = receipt_rmb + tax_refund if receipt_rmb is not None and tax_refund is not None else None
    gross_margin = gross_profit / denominator if gross_profit is not None and denominator not in (None, 0) else None
    employee = orders[0].employee if all(order.employee == orders[0].employee for order in orders) else None
    category = orders[0].category if all(order.category == orders[0].category for order in orders) else None
    commission_sum = complete_sum([order.commission for order in orders])
    commission = round_money(commission_sum) if commission_sum is not None else None
    return Order(client_label or orders[0].client, invoice_label, None, receipt_usd, exchange_rate, receipt_rmb, field_sum("invoiced"), field_sum("actual_cost"), tax_refund, gross_profit, gross_margin, employee, category, commission)


def write_report_sheet(
    sheet,
    orders: list[Order],
    title: str,
    subtitle: str,
    totals: list[Order],
    grand_total: Order | None = None,
    *,
    include_date: bool = True,
    include_commission: bool = False,
) -> None:
    sheet.sheet_view.showGridLines = False
    base_column_count = 10 if include_date else 9
    commission_columns = 3 if include_commission else 0
    last_column = get_column_letter(base_column_count + commission_columns)
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color="17365D")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")

    header_row = 4
    headers = list(TOTAL_HEADERS if include_date else OVERVIEW_HEADERS)
    if include_commission:
        headers.extend(COMMISSION_HEADERS)
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(style="thin", color="B7C9D6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    grand_fill = PatternFill("solid", fgColor="A9D18E")

    def write_order(row: int, order: Order, fill, bold: bool = False) -> None:
        if include_date:
            values = [order.client, order.order_date, order.receipt_usd, order.exchange_rate, order.receipt_rmb, order.invoiced, order.actual_cost, order.tax_refund, order.gross_profit, order.gross_margin]
        else:
            values = [order.client, order.receipt_usd, order.exchange_rate, order.receipt_rmb, order.invoiced, order.actual_cost, order.tax_refund, order.gross_profit, order.gross_margin]
        if include_commission:
            values.extend([order.employee, order.category, order.commission])
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.fill = fill
            cell.border = border
            cell.font = Font(name="Arial", size=10, bold=bold, color="000000")
            cell.alignment = Alignment(vertical="center")
        date_offset = 1 if include_date else 0
        if include_date:
            sheet.cell(row=row, column=2).number_format = "yyyy-mm-dd"
        sheet.cell(row=row, column=2 + date_offset).number_format = '#,##0.00;[Red](#,##0.00);-'
        sheet.cell(row=row, column=3 + date_offset).number_format = "0.0000"
        for column in range(4 + date_offset, 9 + date_offset):
            sheet.cell(row=row, column=column).number_format = '#,##0.00;[Red]-#,##0.00;-'
        sheet.cell(row=row, column=9 + date_offset).number_format = "0.0%"
        if include_commission:
            sheet.cell(row=row, column=base_column_count + 3).number_format = '#,##0.00;[Red]-#,##0.00;-'

    row = 5
    for order in orders:
        write_order(row, order, body_fill)
        row += 1
    if orders and totals:
        row += 1
    for total in totals:
        write_order(row, total, total_fill, True)
        row += 1
    if grand_total is not None:
        row += 1
        write_order(row, grand_total, grand_fill, True)

    if include_date:
        widths = {"A": 18, "B": 14, "C": 15, "D": 12, "E": 16, "F": 16, "G": 16, "H": 15, "I": 17, "J": 15}
    else:
        widths = {"A": 18, "B": 15, "C": 12, "D": 16, "E": 16, "F": 16, "G": 15, "H": 17, "I": 15}
    if include_commission:
        for index, width in enumerate((17, 13, 19), base_column_count + 1):
            widths[get_column_letter(index)] = width
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"


def safe_sheet_name(client: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", client).strip().strip("'") or "客户"
    base = base[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        marker = f"_{suffix}"
        candidate = f"{base[:31 - len(marker)]}{marker}"
        suffix += 1
    used.add(candidate)
    return candidate


def write_employee_sheet(sheet, orders: list[Order]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "员工提成试算"
    sheet["A1"].font = Font(name="Arial", size=15, bold=True, color="17365D")
    sheet["A2"] = "仅2027年适用所附制度；其他年度是规则演示。K类保护期、出货/收汇及费用仍须核实。"
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")
    headers = ["员工姓名", "年度", "客户数", "单据数", "毛利润合计RMB", "逐单提成试算合计RMB", "年度规则后试算RMB"]
    for col, header in enumerate(headers, 1):
        cell = sheet.cell(4, col, header)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    grouped: OrderedDict[tuple[str, int | None], list[Order]] = OrderedDict()
    for order in orders:
        year = order.order_date.year if order.order_date else None
        grouped.setdefault((order.employee or "未填写员工", year), []).append(order)
    undated_employees = {order.employee or "未填写员工" for order in orders if order.order_date is None}
    row = 5
    for (employee, year), group in sorted(grouped.items(), key=lambda item: (item[0][0], item[0][1] or 0)):
        gross_profit = complete_sum([order.gross_profit for order in group])
        commission_sum = complete_sum([order.commission for order in group])
        commission = round_money(commission_sum) if commission_sum is not None else None
        annual = None
        if year == 2027 and employee != "未填写员工" and employee not in undated_employees and gross_profit is not None and commission is not None:
            annual = 0.0 if gross_profit < 0 else commission
        values = [employee, year, len({customer_key(order.client) for order in group}), len(group), gross_profit, commission, annual]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row, col, value)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(name="Arial", size=10, color="000000")
            cell.alignment = Alignment(vertical="center")
        for col in (5, 6, 7):
            sheet.cell(row, col).number_format = '#,##0.00;[Red]-#,##0.00;-'
        row += 1
    for col, width in enumerate((20, 12, 12, 12, 22, 27, 27), 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = "A5"


def build_total_workbook(orders: list[Order], *, include_commission: bool = False) -> tuple[Workbook, int]:
    grouped: OrderedDict[str, list[Order]] = OrderedDict()
    for order in orders:
        grouped.setdefault(order.client, []).append(order)

    workbook = Workbook()
    overview = workbook.active
    overview.title = OVERVIEW_SHEET
    client_totals = [aggregate(client_orders) for client_orders in grouped.values()]
    grand_total = aggregate(orders, client_label="所有客户", invoice_label="所有客户合计")
    write_report_sheet(
        overview,
        [],
        "所有客户总体数据",
        "仅显示各客户合计；所有客户合计位于底部。",
        client_totals,
        grand_total,
        include_date=False,
        include_commission=include_commission,
    )

    if include_commission:
        employee_sheet = workbook.create_sheet("员工提成")
        write_employee_sheet(employee_sheet, orders)

    used_names = {OVERVIEW_SHEET, "员工提成"}
    for client, client_orders in grouped.items():
        sheet = workbook.create_sheet(safe_sheet_name(client, used_names))
        write_report_sheet(
            sheet,
            client_orders,
            f"{client} 收汇及利润汇总",
            "逐单利润取自明细；开票金额为同一票各开票金额之和；缺失数据保持空白。",
            [aggregate(client_orders)],
            include_commission=include_commission,
        )
    return workbook, len(grouped)


def save_atomic(workbook, path: Path) -> None:
    with tempfile.NamedTemporaryFile(prefix=f".{path.stem}_", suffix=path.suffix, dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def process(input_path: Path, output_path: Path, ck_path: Path | None = None) -> tuple[int, int, list[str]]:
    if not input_path.exists() or input_path.suffix.lower() != ".xlsx":
        raise DataFormatError(f"找不到有效的.xlsx文件：{input_path}")
    if input_path.resolve() == output_path.resolve():
        raise DataFormatError("总表必须是一个独立文件，输出路径不能与明细文件相同。")
    if ck_path is not None and ck_path.resolve() == output_path.resolve():
        raise DataFormatError("输出路径不能覆盖C/K标记表。")

    source_workbook = load_workbook(input_path, data_only=False)
    cached_workbook = load_workbook(input_path, data_only=True)
    reader = ValueReader(source_workbook, cached_workbook)
    warnings: list[str] = []
    orders: list[Order] = []
    for sheet in source_workbook.worksheets:
        if sheet.title in TOTAL_SHEETS:
            continue
        blocks = find_blocks(sheet)
        if not blocks:
            warnings.append(f"工作表“{sheet.title}”没有标准明细表头，已跳过。")
            continue
        orders.extend(read_block(sheet, block, reader, warnings) for block in blocks)
    if not orders:
        raise DataFormatError("没有找到可以汇总的客户单据。")

    if ck_path is not None:
        mapping = read_ck_mapping(ck_path)
        apply_commissions(orders, mapping, warnings)
        missing_employees = sorted({order.client for order in orders if order.category is not None and not order.employee})
        if missing_employees:
            warnings.append("下列客户未填员工姓名，无法归入员工年度试算：" + "、".join(missing_employees))

    output_workbook, client_count = build_total_workbook(orders, include_commission=ck_path is not None)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    save_atomic(output_workbook, output_path)

    check = load_workbook(output_path, data_only=True, read_only=True)
    expected_sheets = client_count + (2 if ck_path is not None else 1)
    if OVERVIEW_SHEET not in check.sheetnames or len(check.sheetnames) != expected_sheets:
        raise DataFormatError("输出验证失败：总体数据或客户sheet数量不正确。")
    expected_headers = OVERVIEW_HEADERS + (COMMISSION_HEADERS if ck_path is not None else [])
    headers = [check[OVERVIEW_SHEET].cell(row=4, column=column).value for column in range(1, len(expected_headers) + 1)]
    if headers != expected_headers:
        raise DataFormatError("输出验证失败：总表表头顺序不正确。")
    return len(orders), client_count, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description="读取明细中已有的每单利润，生成总体、客户和可选员工提成sheet。")
    parser.add_argument("input", type=Path, help="已经运行过明细利润程序的.xlsx工作簿")
    parser.add_argument("ck_mapping", nargs="?", type=Path, help="可选：客户名/标记/员工名的C/K标记.xlsx")
    parser.add_argument("-o", "--output", type=Path, help="总表输出路径；默认在明细文件旁生成名字带_总表的独立文件")
    args = parser.parse_args()
    output = args.output or args.input.with_name(f"{args.input.stem}_总表.xlsx")
    try:
        count, client_count, warnings = process(args.input, output, args.ck_mapping)
    except (DataFormatError, PermissionError, OSError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2
    print(f"完成：已生成独立总表，汇总 {client_count} 个客户、{count} 个单据。")
    print(f"文件：{output}")
    if warnings:
        print("需要检查：")
        for warning in warnings:
            print(f"- {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
